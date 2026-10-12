import io
import json
import sys

import pymupdf as fitz
from docx import Document

from app.core.config import settings


def parse(data, mime):
    limit = settings()
    pages = []
    if mime == "application/pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass or document.page_count > limit.max_pages:
                raise ValueError("PDF_PAGE_LIMIT_OR_ENCRYPTED")
            for index, page in enumerate(document):
                text = page.get_text()
                pages.append({"page": index + 1, "text": text, "reference_type": "page"})
                if sum(len(p["text"]) for p in pages) > limit.max_text_chars:
                    raise ValueError("TEXT_LIMIT")
    else:
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        paragraphs += [
            " | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows
        ]
        # DOCX has no stable physical page numbers; use reproducible sections.
        for index in range(0, len(paragraphs), 20):
            pages.append(
                {
                    "page": index // 20 + 1,
                    "text": "\n".join(paragraphs[index : index + 20]),
                    "reference_type": "section",
                }
            )
    if not pages or not any(p["text"].strip() for p in pages):
        raise ValueError("SCANNED_OR_EMPTY_OCR_REQUIRED")
    if sum(len(p["text"]) for p in pages) > limit.max_text_chars or len(pages) > limit.max_pages:
        raise ValueError("TEXT_LIMIT")
    return pages


if __name__ == "__main__":
    try:
        data = sys.stdin.buffer.read(settings().max_upload_bytes + 1)
        output = {"pages": parse(data, sys.argv[1])}
    except Exception as error:
        output = {"error": str(error) if isinstance(error, ValueError) else "PARSE_FAILED"}
    sys.stdout.buffer.write(json.dumps(output).encode())
